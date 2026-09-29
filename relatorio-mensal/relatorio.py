"""Análise financeira mensal da Agência do Futuro.

Lê a planilha "Futuro 2026" (Google Sheets), calcula receita, despesa, margem,
riscos e melhorias, gera dois painéis HTML (vendas e retiradas do sócio) e
envia o resumo + os painéis por um bot do Telegram.

Só lê a planilha. Nunca escreve nela.

Variáveis de ambiente:
  GOOGLE_SHEET_ID               id da planilha (o mesmo do bot)
  GOOGLE_CREDENTIALS_JSON       JSON da conta de serviço (o mesmo do bot)
  TELEGRAM_BOT_TOKEN            token do bot
  TELEGRAM_CHAT_ID              chat que recebe o relatório
  REF_MES                       opcional: último mês analisado (1-12). Padrão: mês
                                anterior se hoje é dia 1 a 3, senão o mês atual (parcial)
  LOCAL_XLSX                    opcional: usa um .xlsx local em vez do Google (teste)
  DRY_RUN=1                     opcional: não envia nada, só grava os arquivos em OUT_DIR
  OUT_DIR                       pasta de saída (padrão: ./saida)
"""
import calendar
import collections
import datetime as dt
import difflib
import io
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

import oportunidades

AQUI = Path(__file__).parent
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
MESES_LONGO = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto',
               'setembro', 'outubro', 'novembro', 'dezembro']
VENDEDORES = {'BRUNO', 'MARIANA'}


# ---------- formatação ----------
def brl(v):
    s = f'{abs(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return ('−' if v < 0 else '') + 'R$ ' + s


def pct(v, casas=1):
    return f'{v * 100:.{casas}f}'.replace('.', ',') + '%'


def num(x):
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace('R$', '').replace(' ', '').strip()
    if not s or s == '-':
        return 0.0
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'\d{1,3}(\.\d{3})+', s):
        s = s.replace('.', '')
    try:
        return float(s)
    except ValueError:
        return 0.0


def txt(x):
    return '' if x is None else str(x).strip()


def data_br(d):
    return d.strftime('%d/%m/%Y')


# ---------- leitura ----------
def baixar_planilha():
    if os.environ.get('LOCAL_XLSX'):
        return Path(os.environ['LOCAL_XLSX']).read_bytes()
    import google.auth.transport.requests
    import requests
    from google.oauth2 import service_account
    info = json.loads(os.environ.get('GOOGLE_CREDENTIALS_JSON') or os.environ['GOOGLE_SERVICE_ACCOUNT_JSON'])
    sheet = os.environ.get('GOOGLE_SHEET_ID') or os.environ['SHEET_ID']
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=['https://www.googleapis.com/auth/drive.readonly',
                      'https://www.googleapis.com/auth/spreadsheets.readonly'])
    cred.refresh(google.auth.transport.requests.Request())
    auth = {'Authorization': f'Bearer {cred.token}'}
    mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    r = requests.get(f'https://www.googleapis.com/drive/v3/files/{sheet}/export',
                     params={'mimeType': mime}, headers=auth, timeout=120)
    if r.status_code == 403:  # Drive API desativada no projeto: tenta a exportação direta do Sheets
        r = requests.get(f'https://docs.google.com/spreadsheets/d/{sheet}/export',
                         params={'format': 'xlsx'}, headers=auth, timeout=120)
    r.raise_for_status()
    return r.content


def cabecalho(ws, linha):
    return {txt(c.value).lower(): c.column for c in ws[linha] if txt(c.value)}


def col(h, *nomes):
    for n in nomes:
        if n.lower() in h:
            return h[n.lower()]
    raise KeyError(f'coluna não encontrada: {nomes}')


def ler_vendas(wv, wf):
    ws, wsf = wv['Vendas'], wf['Vendas']
    h = cabecalho(ws, 1)
    c = dict(mes=col(h, 'Mês'), conta=col(h, 'Conta'), cia=col(h, 'Cia aérea'), cpf=col(h, 'CPF'),
             nome=col(h, 'Nome'), valor=col(h, 'Valor'), base=col(h, 'Base do imposto', 'Lucro na operação'),
             naotrib=col(h, 'Lucro não tributado'), vend=col(h, 'Vendedor'), com=col(h, 'Comissão'),
             pag=col(h, 'Pagamento'), data=col(h, 'Data de venda'),
             ida=col(h, 'Itinerario da ida'), volta=col(h, 'Itinerário da volta'), data_ida=col(h, 'Data da ida'))
    tx, r = [], 2
    while r <= ws.max_row:
        a, b = txt(ws.cell(r, 1).value).lower(), txt(ws.cell(r, 2).value).lower()
        if b == 'total' or a == 'renda mensal' or a == 'mes':
            break
        if any(ws.cell(r, k).value not in (None, '') for k in (c['valor'], c['nome'], c['mes'])):
            d = ws.cell(r, c['data']).value
            tx.append(dict(
                linha=r, mes=int(num(ws.cell(r, c['mes']).value)),
                data=data_br(d) if isinstance(d, dt.datetime) else txt(d),
                conta=txt(ws.cell(r, c['conta']).value), cia=txt(ws.cell(r, c['cia']).value),
                cpf=txt(ws.cell(r, c['cpf']).value), nome=txt(ws.cell(r, c['nome']).value),
                valor=num(ws.cell(r, c['valor']).value), base=num(ws.cell(r, c['base']).value),
                naotrib=num(ws.cell(r, c['naotrib']).value), vend=txt(ws.cell(r, c['vend']).value),
                com=num(ws.cell(r, c['com']).value), pag=txt(ws.cell(r, c['pag']).value),
                ida=txt(ws.cell(r, c['ida']).value), volta=txt(ws.cell(r, c['volta']).value),
                data_ida=ws.cell(r, c['data_ida']).value if isinstance(ws.cell(r, c['data_ida']).value, dt.datetime) else None))
        r += 1
    # quadro-resumo mensal (custos fixos são digitados à mão lá)
    hr = next(i for i in range(r, ws.max_row + 1) if txt(ws.cell(i, 1).value).lower() == 'mes')
    hq = cabecalho(ws, hr)
    q = dict(vendas=col(hq, 'Total vendas'), sal=col(hq, 'Salario'), imp=col(hq, 'Imposto'),
             cart=col(hq, 'Taxas de cartao'), contas=col(hq, 'contas (tel,etc)'), sist=col(hq, 'Sistema'),
             pdd=col(hq, 'PDD e Justiça'))
    resumo = {}
    for i in range(1, 13):
        rr = hr + 1 + i
        resumo[i] = {k: num(ws.cell(rr, v).value) for k, v in q.items()}
        resumo[i]['_data'] = ws.cell(rr, 1).value
        resumo[i]['_sal_formula'] = txt(wsf.cell(rr, q['sal']).value)
    total_row = {k: num(ws.cell(hr + 1, v).value) for k, v in q.items()}
    total_row['_rotulo'] = txt(ws.cell(hr + 1, 1).value)
    renda = dict(valor=num(ws.cell(hr - 1, 2).value), formula=txt(wsf.cell(hr - 1, 2).value),
                 rotulo=txt(ws.cell(hr - 1, 1).value))
    return tx, resumo, total_row, renda


def ler_fluxo(wv):
    if 'Fluxo de Caixa' not in wv.sheetnames:
        return [], []
    ws = wv['Fluxo de Caixa']
    ok, pend = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = row[1] if len(row) > 1 else None
        if not isinstance(d, dt.datetime):
            continue
        v = row[4] if len(row) > 4 else None
        item = dict(data=d.date(), nome=txt(row[3]), valor=v, obs=txt(row[5] if len(row) > 5 else ''))
        (ok if isinstance(v, (int, float)) else pend).append(item)
    return ok, pend


def ler_retiradas(wv, wf):
    if 'Passagens Retiradas' not in wv.sheetnames:
        return None
    ws, wsf = wv['Passagens Retiradas'], wf['Passagens Retiradas']
    h = cabecalho(ws, 1)
    cv, cn = col(h, 'Valor'), col(h, 'Nome')
    totais, carry, formula_ano, linhas_total = {}, None, '', {}
    for r in range(2, ws.max_row + 1):
        a = txt(ws.cell(r, 1).value)
        al = a.lower()
        if al.startswith('total '):
            nome = al.split()[1]
            if nome in MESES_LONGO:
                f = txt(wsf.cell(r, cv).value)
                m = re.search(r'SUM\(N(\d+):N(\d+)\)', f)
                base = re.match(r'=\s*(\d+(?:\.\d+)?)', f)
                i = MESES_LONGO.index(nome) + 1
                totais[i] = dict(linha=r, lo=int(m.group(1)) if m else r, hi=int(m.group(2)) if m else r - 1,
                                 fixo=float(base.group(1)) if base else 0.0, valor=num(ws.cell(r, cv).value))
                linhas_total[r] = i
            elif re.fullmatch(r'total \d{4}', al):
                if wsf.cell(r, 2).value and str(wsf.cell(r, 2).value).startswith('='):
                    formula_ano = str(wsf.cell(r, 2).value)
                else:
                    carry = dict(rotulo=a, valor=num(ws.cell(r, 2).value))
    itens = []
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, dt.datetime) and ws.cell(r, cv).value not in (None, ''):
            itens.append(dict(linha=r, data=a.date(), nome=txt(ws.cell(r, cn).value),
                              rota=txt(ws.cell(r, 6).value), cia=txt(ws.cell(r, 10).value),
                              loc=txt(ws.cell(r, 11).value), valor=num(ws.cell(r, cv).value)))
    return dict(totais=totais, carry=carry, formula_ano=formula_ano, linhas_total=linhas_total, itens=itens)


# ---------- análise ----------
def mes_referencia(hoje):
    if os.environ.get('REF_MES'):
        return int(os.environ['REF_MES']), False
    if hoje.day <= 3:
        return (hoje.month - 1) or 12, False
    ultimo = calendar.monthrange(hoje.year, hoje.month)[1]
    return hoje.month, hoje.day < ultimo


def agrupa(tx, chave, normaliza=None):
    d = collections.defaultdict(lambda: [0.0, 0.0])
    for t in tx:
        k = t[chave].upper()
        if normaliza:
            k = normaliza.get(k, k)
        d[k][0] += t['valor']
        d[k][1] += t['base']
    return d


def unifica_nomes(tx):
    """Junta grafias quase iguais da coluna Conta (ex.: Ancoradouro/Ancoradoouro)."""
    tot = agrupa(tx, 'conta')
    ordem = sorted(tot, key=lambda k: -tot[k][1])
    mapa, canon = {}, []
    for k in ordem:
        alvo = next((c for c in canon if difflib.SequenceMatcher(None, k, c).ratio() >= 0.9
                     or sorted(k.split()) == sorted(c.split())), None)
        if alvo and k:
            mapa[k] = alvo
        else:
            canon.append(k)
    return mapa


def analisar(tx, resumo, total_row, renda, fluxo, pend_fluxo, ret, hoje):
    ref, parcial = mes_referencia(hoje)
    meses = [m for m in range(1, ref + 1) if any(t['mes'] == m for t in tx)]
    ano = hoje.year if ref <= hoje.month else hoje.year - 1
    txp = [t for t in tx if t['mes'] in meses]
    vend = [t for t in txp if t['valor'] > 0]
    M = []
    for m in meses:
        g = [t for t in txp if t['mes'] == m]
        q = resumo[m]
        d = dict(m=m, label=f'{MESES[m - 1]}/{str(ano)[2:]}', vendas=sum(t['valor'] for t in g),
                 trib=sum(t['base'] for t in g), naotrib=sum(t['naotrib'] for t in g),
                 comB=sum(t['com'] for t in g if t['vend'].upper() == 'BRUNO'),
                 comM=sum(t['com'] for t in g if t['vend'].upper() == 'MARIANA'),
                 sal=q['sal'], imp=q['imp'], cart=q['cart'], contas=q['contas'], sist=q['sist'], pdd=q['pdd'])
        d['com'] = d['comB'] + d['comM']
        d['fixo'] = d['sal'] + d['imp'] + d['cart'] + d['contas'] + d['sist'] + d['pdd']
        d['receita'] = d['trib'] + d['naotrib'] + d['com']
        d['despesa'] = d['com'] + d['fixo']
        d['resultado'] = d['trib'] + d['naotrib'] - d['fixo']
        M.append(d)
    S = lambda k: sum(r[k] for r in M)
    ini, fim = M[0], M[-1]
    T = S('trib')

    faixas = []
    for lo, hi, lab in [(0, 5000, 'Até R$ 5 mil'), (5000, 20000, 'R$ 5 mil a 20 mil'), (20000, 1e15, 'Acima de R$ 20 mil')]:
        g = [t for t in vend if lo <= t['valor'] < hi]
        V, L = sum(t['valor'] for t in g), sum(t['base'] for t in g)
        faixas.append(dict(faixa=lab, n=len(g), vendas=V, lucro=L, margem=L / V if V else 0))
    grandes = faixas[2]

    mapa = unifica_nomes(txp)
    tc = sorted(agrupa(txp, 'conta', mapa).items(), key=lambda z: -z[1][1])[:5]
    tp = sorted(agrupa(txp, 'cpf').items(), key=lambda z: -z[1][1])[:5]
    top_cli, top_conta = tp[0], tc[0]

    baixa = [t for t in vend if t['base'] / t['valor'] < 0.04]
    ganho4 = sum(0.04 * t['valor'] - t['base'] for t in baixa if t['base'] >= 0)
    cartao = [t for t in vend if 'cart' in t['pag'].lower()]
    pix = [t for t in vend if 'pix' in t['pag'].lower() and 'cart' not in t['pag'].lower()]
    mg = lambda g: sum(t['base'] for t in g) / sum(t['valor'] for t in g) if g else 0
    Vc = sum(t['valor'] for t in cartao)
    sem_pag = [t for t in vend if not t['pag']]
    pdd = S('pdd')
    res, rec, fixo, com = S('resultado'), S('receita'), S('fixo'), S('com')
    pessoas = S('sal') + com
    rotulo_fim = fim['label'] + (' parcial' if parcial else '')

    # ----- riscos (texto para o dono) -----
    riscos = []
    riscos.append((f'lucro do mês caiu de {brl(ini["resultado"])} ({ini["label"]}) para {brl(fim["resultado"])} ({rotulo_fim}); vendas acima de R$ 20 mil deixam só {pct(grandes["margem"])}',
        'Você está ganhando menos por venda',
        f'O lucro do mês foi de {brl(ini["resultado"])} em {ini["label"]} para {brl(fim["resultado"])} em {rotulo_fim}. '
        f'O quanto a agência fica de cada venda caiu de {pct(ini["trib"] / ini["vendas"])} para {pct(fim["trib"] / fim["vendas"])}. '
        f'As vendas grandes são o problema: as {grandes["n"]} acima de R$ 20 mil somam {pct(grandes["vendas"] / S("vendas"), 0)} do que você vendeu, '
        f'mas dão só {pct(grandes["lucro"] / T, 0)} do lucro (margem de {pct(grandes["margem"])}).'))
    riscos.append((f'custo fixo foi de {pct(ini["fixo"] / ini["receita"], 0)} para {pct(fim["fixo"] / fim["receita"], 0)} da receita do mês',
        'O custo fixo sobe e a receita não acompanha',
        f'O custo fixo foi de {brl(ini["fixo"])} ({ini["label"]}) para {brl(fim["fixo"])} ({rotulo_fim}). '
        f'Hoje ele come {pct(fim["fixo"] / fim["receita"], 0)} do que a agência ganha no mês, contra {pct(ini["fixo"] / ini["receita"], 0)} em {ini["label"]}. '
        f'Salários e comissões são {pct(pessoas / (fixo + com), 0)} de todos os custos, então o lucro depende quase só de vender mais com a mesma equipe.'))
    meses_sem = sorted({MESES[t['mes'] - 1] for t in sem_pag})
    if sem_pag or pdd:
        riscos.append(((f'{len(sem_pag)} vendas sem pagamento anotado ({brl(sum(t["valor"] for t in sem_pag))})' if sem_pag else f'{brl(pdd)} de PDD e Justiça no período'),
            'Dinheiro a receber',
            (f'{len(sem_pag)} vendas ({brl(sum(t["valor"] for t in sem_pag))}, de {", ".join(meses_sem)}) estão sem forma de pagamento anotada.' if sem_pag else '')
            + (f' Já houve {brl(pdd)} de PDD e Justiça.' if pdd else '')))
    riscos.append((f'maior cliente = {pct(top_cli[1][1] / T)} do lucro; conta "{top_conta[0].title()}" = {pct(top_conta[1][1] / T)}',
        'Dependência de um cliente e de uma conta de emissão',
        f'O maior cliente ({"CNPJ " if "/" in top_cli[0] else "CPF "}{mascara(top_cli[0])}) gera {pct(top_cli[1][1] / T)} do lucro. '
        f'A conta de emissão "{top_conta[0].title()}" gera {pct(top_conta[1][1] / T)} do lucro e {pct(top_conta[1][0] / S("vendas"), 0)} de tudo que você vende. '
        'Se essa conta travar, você perde de uma vez a sua maior fonte de emissão.'))
    riscos = [dict(titulo=f'{i + 1}. {r[1]}', texto=r[2], curto=r[0]) for i, r in enumerate(riscos)]

    # ----- melhorias -----
    contas_delta = (fim['contas'] + fim['sist']) - (ini['contas'] + ini['sist'])
    melhorias = []
    if sem_pag:
        melhorias.append(dict(o=f'Cobrar e anotar as {len(sem_pag)} vendas sem forma de pagamento',
            p='Enquanto não estiver anotado, você não sabe se esse dinheiro entrou. É o que mais pesa agora.',
            i=f'{brl(sum(t["valor"] for t in sem_pag))} em risco', t='real', e='Baixo', v=sum(t['valor'] for t in sem_pag)))
    melhorias.append(dict(o='Definir uma margem mínima para vendas acima de R$ 5 mil (por exemplo, 4%)',
        p=f'Venda grande dá trabalho e risco (cartão, parcelamento) e hoje deixa muito pouco. {len(baixa)} vendas ({brl(sum(t["valor"] for t in baixa))}) ficaram abaixo de 4%.',
        i=f'+{brl(ganho4)} no período, se essas vendas tivessem 4%', t='est', e='Baixo', v=ganho4))
    if cartao and pix and mg(cartao) < mg(pix):
        melhorias.append(dict(o='Rever o preço no cartão',
            p=f'No cartão você fica com {pct(mg(cartao), 2)} da venda; no pix, com {pct(mg(pix), 2)}. Se a taxa ou os juros do parcelamento estão saindo do seu bolso, repasse ao cliente ou dê desconto no pix.',
            i=f'+{brl(Vc * 0.01)} para cada 1 ponto percentual a mais no cartão', t='est', e='Médio', v=Vc * 0.01))
    melhorias.append(dict(o=f'Diminuir a dependência do maior cliente e da conta "{top_conta[0].title()}"',
        p='Contrato com prazo e condições para o cliente principal. Ter uma segunda conta ou fornecedor pronto para emitir as mesmas rotas.',
        i=f'{brl(top_cli[1][1])} (cliente) e {brl(top_conta[1][1])} (conta) de lucro exposto', t='real', e='Médio', v=top_cli[1][1]))
    if contas_delta > 0:
        melhorias.append(dict(o='Revisar as contas fixas pequenas ("contas" e "sistema")',
            p=f'Subiram {brl(contas_delta)}/mês desde {ini["label"]}. Confira se cada uma ainda se paga.',
            i=f'Até {brl(contas_delta * 12)}/ano, se voltarem ao nível de {ini["label"]}', t='est', e='Baixo', v=contas_delta * 12))

    # ----- problemas na planilha -----
    problemas = []
    for m in M:
        dif = m['vendas'] - resumo[m['m']]['vendas']
        if abs(dif) > 0.5:
            problemas.append(f'"Total vendas" de {m["label"]} na planilha é {brl(resumo[m["m"]]["vendas"])}, mas a soma dos lançamentos dá {brl(m["vendas"])} (diferença de {brl(dif)}). Provavelmente a fórmula não pega todas as linhas. O painel usa a soma dos lançamentos.')
    d0 = resumo[1]['_data']
    if isinstance(d0, dt.datetime) and d0.year != ano:
        problemas.append(f'O quadro-resumo está rotulado {d0.year} ("{total_row["_rotulo"]}"), mas soma os lançamentos de {ano}.')
    if '/12' in renda['formula'].replace(' ', '') and len(meses) < 12:
        problemas.append(f'"Renda mensal" ({brl(renda["valor"])}) divide o lucro por 12, mas só há {len(meses)} meses. A média real é {brl(res / len(meses))}/mês.')
    soma_sist = sum(resumo[m]['sist'] for m in range(1, 13))
    if abs(total_row['sist'] - soma_sist) > 0.5:
        problemas.append(f'O total de "Sistema" aparece como {brl(total_row["sist"])}, mas a soma dos meses é {brl(soma_sist)}. Algum valor foi digitado como texto.')
    if mapa:
        pares = '; '.join(f'"{k.title()}" = "{v.title()}"' for k, v in mapa.items())
        problemas.append(f'Contas de emissão escritas de mais de um jeito (juntei no painel): {pares}.')
    for t in vend:
        if t['vend'].upper() not in VENDEDORES:
            problemas.append(f'Lançamento de {t["data"]} ({brl(t["valor"])}) com vendedor "{t["vend"] or "vazio"}" e lucro de {brl(t["base"])}. Parece incompleto.')

    # ----- a receber -----
    fc = collections.OrderedDict()
    for f in sorted(fluxo, key=lambda f: f['data']):
        if f['data'] >= dt.date(ano, ref, 1):
            k = f'{MESES[f["data"].month - 1]}/{f["data"].year}'
            fc[k] = fc.get(k, 0) + f['valor']

    D = dict(
        titulo='Painel Financeiro · Agência do Futuro',
        sub=f'Período: 01/01/{ano} a {data_br(hoje) if parcial else data_br(dt.date(ano, ref, calendar.monthrange(ano, ref)[1]))}'
            f'{" (" + MESES_LONGO[ref - 1] + " ainda aberto)" if parcial else ""} · Fonte: planilha "Futuro {ano}" (aba Vendas, {len(txp)} lançamentos) · Gerado em {data_br(hoje)}',
        simples=f'Impostos no período: {brl(S("imp"))}, ou {pct(S("imp") / T, 2)} do Lucro Tributado. Regime: Simples Nacional.',
        M=M, B=faixas, T=T,
        tp=[dict(k=k, doc=('CNPJ ' + k) if '/' in k else ('(sem documento)' if k in ('', '-') else 'CPF ' + mascara(k)),
                 vendas=v[0], lucro=v[1], share=v[1] / T) for k, v in tp],
        tc=[dict(k=k.title(), vendas=v[0], lucro=v[1], share=v[1] / T) for k, v in tc],
        kpis=[
            ['Valor vendido', brl(S('vendas')), f'{len(vend)} vendas com valor'],
            ['Receita da agência', brl(rec), f'{pct(rec / S("vendas"))} do valor vendido'],
            ['Lucro do período', brl(res), f'média de {brl(res / len(M))}/mês · {pct(res / rec)} da receita'],
            ['Custo fixo', brl(fixo), f'{ini["label"]} {brl(ini["fixo"])} → {fim["label"]} {brl(fim["fixo"])}'],
            [f'Lucro {ini["label"]} → {fim["label"]}', f'{brl(ini["resultado"])} → {brl(fim["resultado"])}', 'último mês parcial' if parcial else 'meses fechados'],
        ],
        riscos=riscos, melhorias=melhorias, problemas=problemas,
        salarios='Salários: R$ 2.000,00 Bruno + R$ 2.000,00 Mariana + R$ 1.500,00 Fernanda (em mai/2026, R$ 960,00 a mais para a Mariana). Não há outros encargos.',
        pessoas=f'Salários + comissões: {brl(pessoas)}, ou {pct(pessoas / (fixo + com), 0)} de todos os custos ({brl(fixo + com)}).',
        fc=[[k, v] for k, v in fc.items()], fc_total=sum(fc.values()),
        fc_pend=[f'{data_br(p["data"])} · {p["nome"]}: "{p["valor"]}"' for p in pend_fluxo],
        parcial=parcial,
    )
    R = analisar_retiradas(ret, M, ref, ano, hoje, parcial) if ret else None
    O = oportunidades.analisar(vend, M, sem_pag, S, brl, pct, hoje)
    itens_ret = R.pop('_itens') if R else []
    ctx = dict(parcial=parcial, rotulo=rotulo_fim, res=res, top_conta=(top_conta[0], top_conta[1][0]), itens_ret=itens_ret)
    return D, R, O, ctx


def mascara(doc):
    if '/' in doc:
        return doc
    dig = re.sub(r'\D', '', doc)
    return f'***{dig[3:9]}**' if len(dig) >= 9 else doc


def analisar_retiradas(ret, M, ref, ano, hoje, parcial):
    comB = {r['m']: r['comB'] for r in M}
    linhas = []
    for i in sorted(ret['totais']):
        if i > ref:
            continue
        t = ret['totais'][i]
        g = [x for x in ret['itens'] if t['lo'] <= x['linha'] <= t['hi']]
        retirado = sum(x['valor'] for x in g)
        direito = t['fixo'] + comB.get(i, 0)
        linhas.append(dict(label=f'{MESES[i - 1]}/{str(ano)[2:]}', fixo=t['fixo'], com=comB.get(i, 0), direito=direito,
                           retirado=retirado, n=len(g), saldo=direito - retirado,
                           fora=[x for x in g if x['data'].month != i]))
    carry = ret['carry']['valor'] if ret['carry'] else 0.0
    dir_ = sum(x['direito'] for x in linhas)
    retir = sum(x['retirado'] for x in linhas)
    net = dir_ - retir
    fim = carry + net
    media_dir = dir_ / len(linhas)
    media_net = net / len(linhas)
    acum, pts = carry, [carry]
    for x in linhas:
        acum += x['saldo']
        pts.append(round(acum, 2))
    itens = [x for x in ret['itens'] if any(ret['totais'][i]['lo'] <= x['linha'] <= ret['totais'][i]['hi'] for i in ret['totais'] if i <= ref)]
    por_nome = collections.Counter()
    por_tipo = collections.Counter()
    for x in itens:
        por_nome[x['nome'].split()[0] if x['nome'] else '?'] += x['valor']
        por_tipo['Hotel / pontos' if re.search(r'hotel|accor', (x['rota'] + x['cia'] + x['loc']).lower()) else 'Passagem aérea'] += x['valor']
    maiores = sorted(itens, key=lambda x: -x['valor'])[:5]

    problemas = []
    refs = set(int(n) for n in re.findall(r'N(\d+)', ret['formula_ano']))
    for i, t in ret['totais'].items():
        if t['linha'] not in refs and ret['formula_ano']:
            problemas.append(f'A fórmula "Total {ano}" não soma o total de {MESES_LONGO[i - 1]} (linha {t["linha"]}), então deixa de fora {brl(t["valor"])}.')
    for r in refs:
        if r not in ret['linhas_total']:
            problemas.append(f'A fórmula "Total {ano}" soma a célula N{r}, que não é o total de nenhum mês.')
    futuros = sum(t['fixo'] for i, t in ret['totais'].items() if i > ref)
    if futuros:
        problemas.append(f'O saldo da planilha já conta o direito dos meses que ainda não chegaram ({brl(futuros)}) como se estivesse ganho.')
    for x in linhas:
        if x['fora']:
            problemas.append(f'{len(x["fora"])} retiradas lançadas no grupo de {x["label"]} são de outro mês (datas de {data_br(min(f["data"] for f in x["fora"]))} a {data_br(max(f["data"] for f in x["fora"]))}).')
    if ret['carry']:
        problemas.append(f'O saldo trazido ("{ret["carry"]["rotulo"]}", {brl(carry)}) foi digitado como número, sem fórmula. Não dá para conferir de onde veio.')

    leitura = [
        f'<b>No ano, o saldo é {"positivo" if net >= 0 else "negativo"}:</b> você retirou {brl(retir)}, contra um direito de {brl(dir_)} ({brl(net)}).',
        f'<b>Saldo acumulado com a agência:</b> {brl(fim)}. Isso é ' + f'{abs(fim) / media_dir:.1f}'.replace('.', ',') + f' meses do seu direito médio ({brl(media_dir)}/mês).' if fim < 0 else f'<b>Saldo acumulado com a agência:</b> {brl(fim)} a seu favor.',
    ]
    if fim < 0 and media_net > 0:
        leitura.append(f'<span class="tag est">ESTIMATIVA</span> No ritmo atual (sobra média de {brl(media_net)}/mês), zerar esse saldo levaria cerca de {round(abs(fim) / media_net)} meses. A premissa é que comissões e retiradas continuem na média do período.')
    maior = max(linhas, key=lambda x: x['retirado'])
    leitura.append(f'<b>O maior mês</b> foi {maior["label"]}, com {brl(maior["retirado"])} ({pct(maior["retirado"] / retir, 0)} do total retirado).')
    leitura.append(f'<b>Seu direito depende das vendas:</b> {pct(sum(x["com"] for x in linhas) / dir_, 0)} dele vem das suas comissões.')
    leitura.append('<b>Os seus R$ 2.000,00 fixos já estão na linha "Salário" da aba Vendas</b>, e a comissão também. As passagens são só a forma de pagamento, então não há custo em dobro.')

    return dict(
        _itens=itens,
        sub=f'Período: 01/01/{ano} a {data_br(hoje) if parcial else data_br(dt.date(ano, ref, calendar.monthrange(ano, ref)[1]))} · Fonte: aba "Passagens Retiradas" · Gerado em {data_br(hoje)}',
        M=linhas and [{k: v for k, v in x.items() if k != 'fora'} for x in linhas],
        pts=pts, carry=carry, dir=dir_, ret=retir, net=net, fim=fim,
        kpis=[['Direito no período', brl(dir_), f'fixo {brl(sum(x["fixo"] for x in linhas))} + comissões {brl(sum(x["com"] for x in linhas))}'],
              ['Retirado em passagens', brl(retir), f'{len(itens)} lançamentos'],
              ['Saldo do ano', brl(net), 'retirou menos que o direito' if net >= 0 else 'retirou mais que o direito'],
              ['Saldo acumulado', brl(fim), f'inclui {brl(carry)} do ano anterior']],
        nomes=[[k, v] for k, v in por_nome.most_common()], tipos=[[k, v] for k, v in por_tipo.most_common()],
        maiores=[[f'{data_br(x["data"])} · {x["nome"].split()[0]} · {x["rota"] if x["rota"] not in ("", "-") else x["loc"]}', x['valor']] for x in maiores],
        leitura=leitura, problemas=problemas)


# ---------- saída ----------
def render(nome, dados):
    css = (AQUI / 'templates' / 'estilo.css').read_text()
    html = (AQUI / 'templates' / nome).read_text()
    return html.replace('/*__CSS__*/', css).replace('/*__DATA__*/null', json.dumps(dados, ensure_ascii=False, default=str))


def _pedacos(texto, limite=3900):
    """Telegram aceita até 4096 caracteres por mensagem: quebra em parágrafos se passar."""
    if len(texto) <= limite:
        return [texto]
    partes, atual = [], ''
    for par in texto.split('\n\n'):
        if atual and len(atual) + len(par) + 2 > limite:
            partes.append(atual)
            atual = par
        else:
            atual = f'{atual}\n\n{par}' if atual else par
    return partes + [atual]


def enviar_telegram(seq, arquivos):
    """Texto e imagens aparecem em qualquer celular; os HTML vão no fim, como extra."""
    import requests
    token, chat = os.environ['TELEGRAM_BOT_TOKEN'], os.environ['TELEGRAM_CHAT_ID']
    base = f'https://api.telegram.org/bot{token}'
    for tipo, conteudo in seq:
        if tipo == 'texto':
            for parte in _pedacos(conteudo):
                r = requests.post(f'{base}/sendMessage', data=dict(chat_id=chat, text=parte, parse_mode='HTML',
                                                                    disable_web_page_preview='true'), timeout=60)
                r.raise_for_status()
        else:
            caminho, legenda = conteudo
            with open(caminho, 'rb') as f:
                r = requests.post(f'{base}/sendPhoto', data=dict(chat_id=chat, caption=legenda), files=dict(photo=f), timeout=120)
            r.raise_for_status()
    for p in arquivos:
        with open(p, 'rb') as f:
            r = requests.post(f'{base}/sendDocument', data=dict(chat_id=chat, caption='Extra: abra no navegador do computador (no celular o gráfico pode não aparecer).'),
                              files=dict(document=(p.name, f, 'text/html')), timeout=120)
        r.raise_for_status()


def main():
    hoje = dt.date.fromisoformat(os.environ['HOJE']) if os.environ.get('HOJE') else dt.date.today()
    conteudo = baixar_planilha()
    wv = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    wf = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=False)
    tx, resumo, total_row, renda = ler_vendas(wv, wf)
    fluxo, pend = ler_fluxo(wv)
    ret = ler_retiradas(wv, wf)
    D, R, O, ctx = analisar(tx, resumo, total_row, renda, fluxo, pend, ret, hoje)
    out = Path(os.environ.get('OUT_DIR', 'saida'))
    out.mkdir(parents=True, exist_ok=True)
    sufixo = f'{hoje:%Y-%m}'
    arquivos = [out / f'painel-{sufixo}.html']
    arquivos[0].write_text(render('painel.html', D))
    if R:
        arquivos.append(out / f'retiradas-{sufixo}.html')
        arquivos[1].write_text(render('retiradas.html', R))
    I = oportunidades.insights_retiradas(ctx['itens_ret'], R)
    imgs = oportunidades.graficos(O, out, sufixo, I)
    seq = oportunidades.sequencia(O, ctx['parcial'], ctx['rotulo'], ctx['res'], ctx['top_conta'], R, brl, pct, MESES, imgs, I)
    texto = '\n\n----\n\n'.join(c for t, c in seq if t == 'texto')
    (out / f'mensagem-{sufixo}.txt').write_text(texto)
    print(texto)
    if os.environ.get('DRY_RUN') == '1':
        print('\nDRY_RUN: nada enviado. Arquivos em', out.resolve())
        return
    enviar_telegram(seq, arquivos)
    print('\nEnviado para o Telegram.')


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print('ERRO:', repr(e), file=sys.stderr)
        raise
